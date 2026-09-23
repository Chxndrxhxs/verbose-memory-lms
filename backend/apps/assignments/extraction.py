"""Extract readable text and embedded images from an uploaded source document.

PDF / DOCX / plain text are supported. Images extracted from PDF pages are
saved through the document storage backend and returned as media URLs so
they can be embedded into generated questions and options.
"""

import logging
import re
from pathlib import Path

from django.conf import settings

logger = logging.getLogger(__name__)

MAX_EXTRACT_CHARS = 40_000
MAX_PAGE_IMAGES = 8

_TEXT_EXTENSIONS = {".txt", ".md"}

_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}


def document_file_id(url: str) -> str:
    """Stable storage id for an uploaded document (path under MEDIA_ROOT).

    Keeps subdirectory structure (``lessons/<name>``) so the storage backend
    can later move to S3 without changing every caller.
    """
    media_url = str(settings.MEDIA_URL or "media/").lstrip("/")
    clean = str(url).split("?", 1)[0]
    if media_url and media_url in clean:
        name = clean.split(media_url, 1)[1].lstrip("/")
    else:
        name = Path(clean).name
    if not name:
        raise ValueError("Could not resolve document path from source URL")
    return name


def document_path_from_url(url: str) -> Path:
    """Resolve a stored media URL back to its on-disk path under MEDIA_ROOT.

    `source_document` may be a relative URL ("media/lessons/<name>") or an absolute
    one ("http://host/media/lessons/<name>"); the path after the media prefix
    is preserved so subdirectories like ``lessons/`` still resolve.
    """
    return Path(settings.MEDIA_ROOT) / document_file_id(url)


def extract_document_markdown(url: str) -> str:
    """Convert the document at `url` to markdown via MarkItDown.

    Markdown preserves question numbering, option lists, and tables far
    better than raw pypdf text, so it is the preferred input for the
    verbatim Q&A extractor. Raises when MarkItDown is unavailable or fails.
    """
    from markitdown import MarkItDown

    path = document_path_from_url(url)
    if not path.exists():
        raise FileNotFoundError(f"Source document not found on disk: {path}")
    result = MarkItDown().convert(str(path))
    markdown = (result.text_content or "").strip()
    logger.info("MarkItDown extracted %s chars from %s", len(markdown), path.name)
    return markdown[:MAX_EXTRACT_CHARS]


def split_markdown_pages(markdown: str) -> list[str]:
    """Split MarkItDown output into per-page chunks.

    Handles MarkItDown page-break markers as well as printed "Page N of M"
    footers, which is what question-paper PDFs usually contain.
    """
    import re

    text = markdown or ""
    markers = list(
        re.finditer(
            r"(?im)^[ \t]*(?:<!--\s*page\s*\d+.*?-->|---\s*page\s*\d+.*?|"
            r"page\s+\d+\s+of\s+\d+|\\f)\s*$",
            text,
        )
    )
    if len(markers) <= 1:
        chunks = [text] if text.strip() else []
        return chunks
    chunks = []
    for i in range(len(markers)):
        start = markers[i].end()
        end = markers[i + 1].start() if i + 1 < len(markers) else len(text)
        chunk = text[start:end].strip()
        header = text[markers[i].start() : markers[i].end()].strip()
        chunks.append(f"{header}\n{chunk}".strip())
    return chunks


def split_answer_key(markdown: str) -> tuple[str, str]:
    """Split off a trailing answer-key section so it can ride along per page."""
    import re

    match = re.search(r"(?im)^[ \t]*.{0,20}answer\s*key.{0,20}$", markdown or "")
    if not match:
        return markdown, ""
    return markdown[: match.start()].strip(), markdown[match.start() :].strip()


def extract_document_text(url: str) -> str:
    """Return up to MAX_EXTRACT_CHARS of text from the document at `url`."""
    text = "\n".join(page["text"] for page in extract_document_pages(url))
    text = " ".join(text.split())
    logger.info("Extracted %s chars of text from %s", len(text), url)
    return text[:MAX_EXTRACT_CHARS]


def extract_document_pages(url: str) -> list[dict]:
    """Extract per-page content: text plus any embedded images (saved to disk).

    Returns a list of ``{"page": int, "text": str, "images": [{"name", "url",
    "path", "file_id"}]}``. Images go through the document storage backend
    (local ``MEDIA_ROOT`` today, S3-compatible tomorrow); ``url`` keeps the
    MEDIA_URL shape so the client resolves them with absoluteMediaUrl().
    """
    path = document_path_from_url(url)
    if not path.exists():
        raise FileNotFoundError(f"Source document not found on disk: {path}")
    ext = path.suffix.lower()

    if ext == ".pdf":
        pages = _extract_pdf_pages(path)
    elif ext == ".docx":
        pages = _extract_docx_pages(path)
    elif ext in _TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8", errors="ignore")
        pages = [{"page": 1, "text": text, "images": []}]
    else:
        raise ValueError(f"Unsupported document type for extraction: {ext}")

    pages = [
        {
            "page": p["page"],
            "text": " ".join(p.get("text", "").split()),
            "images": p.get("images", []),
        }
        for p in pages
    ]
    logger.info("Extracted %s pages from %s", len(pages), path.name)
    return pages


def save_figure(data: bytes, file_id: str) -> tuple[str, str]:
    """Persist one extracted figure; returns ``(url, path)``.

    Single choke point for figure storage: local disk today, swap the body
    for S3/boto3 later without touching callers.
    """
    disk_path = Path(settings.MEDIA_ROOT) / file_id
    disk_path.parent.mkdir(parents=True, exist_ok=True)
    disk_path.write_bytes(bytes(data))
    url = f"{settings.MEDIA_URL.rstrip('/')}/{file_id}"
    return url, str(disk_path)


def _image_store_dir(path: Path) -> tuple[Path, str]:
    slug = re.sub(r"[^a-zA-Z0-9_-]", "_", path.stem)[:60] or "document"
    target = Path(settings.MEDIA_ROOT) / "assignment_images" / slug
    target.mkdir(parents=True, exist_ok=True)
    return target, slug


def _figure_file_id(slug: str, name: str) -> str:
    return f"assignment_images/{slug}/{name}"


def _extract_pdf_pages(path: Path) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    _target, slug = _image_store_dir(path)
    pages: list[dict] = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        images = []
        try:
            raw_images = list(page.images)
        except Exception:
            logger.warning("Could not read images on PDF page %s", number, exc_info=True)
            raw_images = []
        for index, image in enumerate(raw_images):
            if len(images) >= MAX_PAGE_IMAGES:
                break
            payload = getattr(image, "data", None)
            if not payload:
                continue
            name = f"p{number:02d}_img{index:02d}{_image_suffix(image)}"
            try:
                url, disk_path = save_figure(bytes(payload), _figure_file_id(slug, name))
                images.append(
                    {
                        "name": name,
                        "url": url,
                        "path": disk_path,
                        "file_id": _figure_file_id(slug, name),
                    }
                )
            except Exception:
                logger.warning(
                    "Could not save image %s from PDF page %s", name, number, exc_info=True
                )
        pages.append({"page": number, "text": text, "images": images})
    return pages


def _extract_docx_pages(path: Path) -> list[dict]:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(parts)
    return [{"page": 1, "text": text, "images": []}]


def _image_suffix(image) -> str:
    name = str(getattr(image, "name", "") or "")
    suffix = Path(name).suffix.lower()
    if suffix in _IMAGE_SUFFIXES:
        return suffix
    return ".png"
